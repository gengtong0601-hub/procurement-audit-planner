"""Full-population indicators. All values are computed here, never by a model."""

import csv
from collections import defaultdict, deque
from datetime import date, datetime, timedelta
from decimal import Decimal, InvalidOperation
from pathlib import Path

from openpyxl import load_workbook

from procurement_audit.models import Analytics, Indicator, RuleBook, Transaction
from procurement_audit.storage import WorkflowError, load_manifest, source_hashes, source_path

SPEND_COLUMNS = [
    "transaction_id",
    "supplier",
    "department",
    "category",
    "payment_date",
    "amount",
    "invoice_number",
    "contract_id",
]
CONTRACT_COLUMNS = ["contract_id", "supplier", "category", "start_date", "end_date"]
PARAMETERS = {"threshold_lower_share": "0.90", "split_window_days": 30, "sample_limit": 10}


def as_date(value) -> date:
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    return date.fromisoformat(str(value))


def read_transactions(case: Path) -> list[Transaction]:
    manifest = load_manifest(case)
    rows, seen = [], set()
    for source in manifest.spend:
        workbook = load_workbook(source_path(case, source.file), read_only=True, data_only=True)
        try:
            if source.sheet not in workbook.sheetnames:
                raise WorkflowError(f"Missing sheet {source.sheet} in {source.file}")
            sheet_rows = workbook[source.sheet].iter_rows(values_only=True)
            headers = next(sheet_rows, ())
            names = [str(value).strip() if value is not None else "" for value in headers]
            if len(names) != len(set(names)) or not set(SPEND_COLUMNS) <= set(names):
                raise WorkflowError(
                    f"Invalid spend columns in {source.file}; see docs/data-contract.md."
                )
            for index, values in enumerate(sheet_rows, 2):
                if all(value is None for value in values):
                    continue
                raw = dict(zip(names, values, strict=False))
                try:
                    amount = Decimal(str(raw["amount"]))
                    if not amount.is_finite() or amount != amount.quantize(Decimal("0.01")):
                        raise ValueError("Amounts must be finite with at most two decimals.")
                    required = SPEND_COLUMNS[:5]
                    if any(
                        raw.get(name) is None or not str(raw[name]).strip() for name in required
                    ):
                        raise ValueError("Missing required value.")
                    row = Transaction(
                        **{key: str(raw[key]).strip() for key in required if key != "payment_date"},
                        payment_date=as_date(raw["payment_date"]),
                        amount=amount,
                        invoice_number=str(raw.get("invoice_number") or "").strip(),
                        contract_id=str(raw.get("contract_id") or "").strip(),
                        source_file=source.file,
                        source_sheet=source.sheet,
                        source_row=index,
                    )
                except (ValueError, TypeError, InvalidOperation) as exc:
                    raise WorkflowError(
                        f"Invalid spend record: {source.file}/{source.sheet}, row {index}"
                    ) from exc
                if row.transaction_id in seen:
                    raise WorkflowError(f"Duplicate transaction ID: {row.transaction_id}")
                if not manifest.period_start <= row.payment_date <= manifest.period_end:
                    raise WorkflowError(
                        f"Transaction {row.transaction_id} is outside the audit period."
                    )
                seen.add(row.transaction_id)
                rows.append(row)
        finally:
            workbook.close()
    if not rows:
        raise WorkflowError("No spend rows were found; an empty population cannot be analysed.")
    if len(rows) > 100000:
        raise WorkflowError(
            "This local implementation supports up to 100,000 payment rows per case."
        )
    return rows


def read_contracts(case: Path) -> dict:
    manifest = load_manifest(case)
    contracts = {}
    with source_path(case, manifest.contracts).open(encoding="utf-8-sig", newline="") as file:
        reader = csv.DictReader(file)
        if not set(CONTRACT_COLUMNS) <= set(reader.fieldnames or []):
            raise WorkflowError("Invalid contracts CSV columns; see docs/data-contract.md.")
        for index, raw in enumerate(reader, 2):
            try:
                values = {key: raw[key].strip() for key in CONTRACT_COLUMNS}
                if not all(values.values()) or values["contract_id"] in contracts:
                    raise ValueError("Blank field or duplicate contract ID.")
                start, end = as_date(values["start_date"]), as_date(values["end_date"])
                if end < start:
                    raise ValueError("Contract end precedes its start.")
                contracts[values["contract_id"]] = {**values, "start": start, "end": end}
            except (ValueError, AttributeError, TypeError) as exc:
                raise WorkflowError(f"Invalid contract at CSV row {index}") from exc
    return contracts


def analyse(case: Path, rules: RuleBook) -> Analytics:
    rows = read_transactions(case)
    contracts = read_contracts(case)
    indicators = []
    for rule in rules.rules:
        selected = {}
        if rule.indicator == "threshold_clustering":
            lower = rule.threshold * Decimal(PARAMETERS["threshold_lower_share"])
            selected = {r.transaction_id: r for r in rows if lower <= r.amount < rule.threshold}
            description = f"Payments from {lower} inclusive to {rule.threshold} exclusive."
        elif rule.indicator == "split_purchases":
            groups = defaultdict(list)
            for row in rows:
                if 0 < row.amount < rule.threshold:
                    groups[(row.supplier, row.department, row.category)].append(row)
            for group in groups.values():
                window, total = deque(), Decimal(0)
                for row in sorted(group, key=lambda r: (r.payment_date, r.transaction_id)):
                    while window and (row.payment_date - window[0].payment_date) > timedelta(
                        days=30
                    ):
                        total -= window.popleft().amount
                    window.append(row)
                    total += row.amount
                    if len(window) >= 2 and total >= rule.threshold:
                        selected.update({r.transaction_id: r for r in window})
            description = (
                f"Same supplier/department/category: two or more positive payments below {rule.threshold}, "
                "within an inclusive 30-day window, whose sum reaches the threshold."
            )
        elif rule.indicator == "off_contract":
            for row in rows:
                contract = contracts.get(row.contract_id)
                if not contract or not (
                    contract["supplier"] == row.supplier
                    and contract["category"] == row.category
                    and contract["start"] <= row.payment_date <= contract["end"]
                ):
                    selected[row.transaction_id] = row
            description = (
                "No exact contract-ID, supplier, category and validity-date match in the supplied register; "
                "an incomplete public register is an alternative explanation."
            )
        elif rule.indicator == "duplicates":
            groups = defaultdict(list)
            for row in rows:
                if row.invoice_number and row.amount > 0:
                    groups[(row.supplier, row.invoice_number, row.amount)].append(row)
            for group in groups.values():
                if len(group) > 1:
                    selected.update({row.transaction_id: row for row in group})
            description = "Repeated supplier + nonblank invoice number + positive amount; all group rows included."
        else:
            continue
        indicators.append(
            Indicator(
                indicator_id=f"IND-{rule.rule_id.removeprefix('RULE-')}",
                rule_id=rule.rule_id,
                kind=rule.indicator,
                transaction_ids=sorted(selected),
                gross_value=sum((abs(r.amount) for r in selected.values()), Decimal(0)),
                count=len(selected),
                description=description,
            )
        )
    return Analytics(
        row_count=len(rows),
        net_spend=sum((r.amount for r in rows), Decimal(0)),
        gross_spend=sum((abs(r.amount) for r in rows), Decimal(0)),
        supplier_count=len({r.supplier for r in rows}),
        transactions=rows,
        indicators=indicators,
        parameters=PARAMETERS,
        source_hashes=source_hashes(case),
    )
