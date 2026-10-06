from decimal import Decimal

import pytest
from openpyxl import load_workbook
from pydantic import ValidationError

from procurement_audit.analytics import analyse, read_transactions
from procurement_audit.cli import main
from procurement_audit.decisions import read_decisions
from procurement_audit.demo import create_demo
from procurement_audit.models import Decision, Decisions, Signoff
from procurement_audit.qa import check_numbers, check_samples, check_traceability
from procurement_audit.storage import WorkflowError, read_json, write_json
from procurement_audit.workflow import (
    analysis_artifacts,
    build_pack,
    case_status,
    extract_rules,
    run_analysis,
    save_decisions,
    sign_off,
    verify_pack,
)


def approve(case):
    register = analysis_artifacts(case)[3]
    decisions = Decisions(
        reviewer="Synthetic test auditor",
        decisions=[
            Decision(
                risk_id=r.risk_id, decision="approve", comment="Automated synthetic test only."
            )
            for r in register.risks
        ],
    )
    save_decisions(case, decisions)
    return decisions


def test_full_population_and_four_indicators(case):
    rules, analytics, history, register = analysis_artifacts(case)
    assert analytics.row_count == 12
    assert analytics.net_spend == Decimal("467000")
    assert analytics.gross_spend == Decimal("511000")
    expected = {
        "threshold_clustering": ["PAY-001", "PAY-002"],
        "split_purchases": ["PAY-001", "PAY-002", "PAY-003"],
        "off_contract": ["PAY-006"],
        "duplicates": ["PAY-004", "PAY-005"],
    }
    assert {i.kind: i.transaction_ids for i in analytics.indicators} == expected
    assert analytics.transactions[0].source_row == 2
    assert analytics.transactions[8].amount == Decimal("-22000")
    assert len(rules.rules) == 5 and len(history.items) == 4 and len(register.risks) == 5


def test_human_gate_blocks_planning_and_cli_returns_failure(case, capsys):
    with pytest.raises(WorkflowError, match="Human gate is pending"):
        build_pack(case)
    assert main(["plan", "--case", str(case)]) == 1
    assert "STOP:" in capsys.readouterr().err
    assert not (case / "outputs/planning-memo.md").exists()


def test_complete_pack_and_explicit_signoff(case):
    approve(case)
    program = build_pack(case)
    assert sum(t.days for t in program.tests) == 15
    assert len(program.tests) == 5
    assert all(c.passed for c in verify_pack(case))
    assert "PENDING" in (case / "outputs/review.md").read_text()
    assert case_status(case)["stage"] == "signoff"
    sign_off(
        case,
        Signoff(
            reviewer="Synthetic test auditor",
            decision="sign_off",
            comment="Automated fixture sign-off only.",
        ),
    )
    assert case_status(case)["stage"] == "signed_off"
    assert all(
        (case / "outputs" / name).is_file()
        for name in [
            "planning-memo.md",
            "risk-control-matrix.md",
            "audit-program.md",
            "review.md",
            "team-messages.json",
        ]
    )
    messages = read_json(case / "outputs/team-messages.json")
    assert {m["sender"] for m in messages} == {"team-lead", "planner", "challenger", "qa-reviewer"}


def test_amendment_changes_priority_and_rejection_excludes_risk(case):
    decisions = approve(case)
    decisions.decisions[0] = Decision(
        risk_id="RISK-001", decision="reject", comment="Reviewed test exclusion."
    )
    decisions.decisions[-1] = Decision(
        risk_id="RISK-005",
        decision="amend",
        likelihood=5,
        impact=5,
        comment="Reviewed higher control priority.",
    )
    decisions.decisions[1] = Decision(
        risk_id="RISK-002",
        decision="amend",
        likelihood=2,
        impact=2,
        comment="Synthetic test: corroborating evidence lowers this planning priority.",
    )
    save_decisions(case, decisions)
    program = build_pack(case)
    assert program.tests[0].risk_id == "RISK-005"
    assert "RISK-001" not in {t.risk_id for t in program.tests}
    assert "RULE-001" in program.excluded_rules
    assert "5 × 5" in (case / "outputs/planning-memo.md").read_text()


def test_all_rejected_has_no_invented_fieldwork(case):
    decisions = approve(case)
    for decision in decisions.decisions:
        decision.decision = "reject"
        decision.comment = "Synthetic test: scope rejected."
    save_decisions(case, decisions)
    program = build_pack(case)
    assert program.tests == []
    assert len(program.excluded_rules) == 5
    assert all(check.passed for check in verify_pack(case))


def test_missing_or_duplicate_decisions_are_blocked(case):
    decisions = approve(case)
    decisions.decisions[-1] = decisions.decisions[0]
    with pytest.raises(WorkflowError, match="unique decision"):
        save_decisions(case, decisions)


def test_rejection_and_amendment_require_explanation():
    with pytest.raises(ValidationError):
        Decision(risk_id="RISK-001", decision="reject")
    with pytest.raises(ValidationError):
        Decision(risk_id="RISK-001", decision="amend", comment="Missing ratings.")


def test_false_quote_is_rejected(tmp_path):
    create_demo(tmp_path / "case")
    case = tmp_path / "case"
    candidates = read_json(case / "inputs/rule-candidates.json")
    candidates["rules"][0]["quote"] = "A fabricated quote that does not exist in the rulebook."
    write_json(case / "inputs/rule-candidates.json", candidates)
    with pytest.raises(WorkflowError, match="Unverified quote"):
        extract_rules(case)
    assert not (case / "outputs/rules.json").exists()


def test_invented_threshold_is_rejected(tmp_path):
    case = tmp_path / "case"
    create_demo(case)
    candidates = read_json(case / "inputs/rule-candidates.json")
    candidates["rules"][0]["threshold"] = "50000"
    write_json(case / "inputs/rule-candidates.json", candidates)
    with pytest.raises(WorkflowError, match="does not appear"):
        extract_rules(case)


def test_missing_source_stops_without_guessing(case):
    (case / "inputs/contracts.csv").unlink()
    with pytest.raises(WorkflowError, match="Missing required source"):
        run_analysis(case)


def test_changed_source_blocks_planning_and_signoff(case):
    approve(case)
    build_pack(case)
    with (case / "inputs/contracts.csv").open("a") as file:
        file.write("\n")
    with pytest.raises(WorkflowError, match="Source files changed"):
        build_pack(case)
    with pytest.raises(WorkflowError, match="Source files changed"):
        sign_off(
            case,
            Signoff(reviewer="Test auditor", decision="sign_off", comment="Should be blocked."),
        )
    assert case_status(case)["stale"]


def test_modified_pack_cannot_be_signed_off(case):
    approve(case)
    build_pack(case)
    (case / "outputs/audit-program.md").write_text("Edited after QA.")
    with pytest.raises(WorkflowError, match="Planning artifact changed"):
        verify_pack(case)


def test_analysis_rerun_archives_decisions_and_resets_gate(case):
    approve(case)
    build_pack(case)
    run_analysis(case)
    archived = list((case / "archives").iterdir())
    assert len(archived) == 1
    assert (archived[0] / "planning-memo.md").exists()
    assert not (case / "outputs/planning-memo.md").exists()
    assert all(
        d.decision == "pending"
        for d in read_decisions(case / "outputs/auditor-comments.md").decisions
    )


def test_numbers_and_samples_checks_detect_corruption(case):
    decisions = approve(case)
    program = build_pack(case)
    rules, analytics, history, register = analysis_artifacts(case)
    bad = analytics.model_copy(update={"net_spend": Decimal("1")})
    with pytest.raises(WorkflowError, match="differ"):
        check_numbers(case, rules, bad, register)
    sampled = next(t for t in program.tests if t.sample_ids)
    sampled.sample_locations[0] = "spend.xlsx / Payments / row 999"
    with pytest.raises(WorkflowError, match="Incorrect workbook row"):
        check_samples(case, rules, analytics, register, program)
    program.tests.pop()
    with pytest.raises(WorkflowError, match="every approved risk"):
        check_traceability(rules, analytics, history, register, decisions, program)


@pytest.mark.parametrize("offset,expected", [(30, True), (31, False)])
def test_split_window_boundary(tmp_path, offset, expected):
    from datetime import date, timedelta

    case = tmp_path / "case"
    create_demo(case)
    rules = extract_rules(case)
    workbook = load_workbook(case / "inputs/spend.xlsx")
    sheet = workbook.active
    sheet.delete_rows(2, sheet.max_row)
    for index, day in enumerate([date(2025, 1, 1), date(2025, 1, 1) + timedelta(days=offset)], 1):
        sheet.append(
            [
                f"PAY-{index}",
                "Atlas Works",
                "Estates",
                "Repairs",
                day.isoformat(),
                60000,
                f"INV-{index}",
                "C-001",
            ]
        )
    workbook.save(case / "inputs/spend.xlsx")
    analytics = analyse(case, rules)
    match = next(i for i in analytics.indicators if i.kind == "split_purchases")
    assert bool(match.count) is expected


def test_malformed_spend_rows_are_not_silently_skipped(case):
    workbook = load_workbook(case / "inputs/spend.xlsx")
    workbook.active["F3"] = "not money"
    workbook.save(case / "inputs/spend.xlsx")
    with pytest.raises(WorkflowError, match="row 3"):
        read_transactions(case)


def test_budget_does_not_silently_drop_approved_risks(case):
    manifest = read_json(case / "inputs/manifest.json")
    manifest["auditor_days"] = 1
    write_json(case / "inputs/manifest.json", manifest)
    run_analysis(case)
    approve(case)
    with pytest.raises(WorkflowError, match="cannot cover all approved risks"):
        build_pack(case)


def test_empty_history_candidates_are_not_assumed_to_mean_no_history(case):
    write_json(
        case / "inputs/history-candidates.json", {"items": [], "coverage_notes": ["Not reviewed."]}
    )
    with pytest.raises(WorkflowError, match="Every supplied committee paper"):
        run_analysis(case)


def test_no_credentials_required_for_default_or_invented_for_llm(case, monkeypatch):
    monkeypatch.delenv("AUDIT_LLM_API_KEY", raising=False)
    run_analysis(case)
    with pytest.raises(WorkflowError, match="requires AUDIT_LLM_API_KEY"):
        extract_rules(case, use_llm=True)
